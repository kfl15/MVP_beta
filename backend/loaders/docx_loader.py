from typing import List
from docx import Document


def load_docx(file_path: str) -> str:
    """
    Read a .docx file and return plain text.
    Keeps paragraph breaks.
    """
    doc = Document(file_path)
    lines: List[str] = []

    for p in doc.paragraphs:
        text = (p.text or "").strip()
        if text:
            lines.append(text)

    # Tables (optional but useful)
    for table in doc.tables:
        for row in table.rows:
            row_text = []
            for cell in row.cells:
                cell_text = (cell.text or "").strip()
                if cell_text:
                    row_text.append(cell_text)
            if row_text:
                lines.append(" | ".join(row_text))

    return "\n".join(lines).strip()
